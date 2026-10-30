"""xlsx → txt 转换测试：内容保真、结构不乱。

运行：python tests/test_xlsx_convert.py
"""
import datetime
import os
import sys
import unittest
from io import BytesIO

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from services.file_ops import xlsx_to_lines  # noqa: E402


def _make_xlsx(rows: list[list]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestXlsxConvert(unittest.TestCase):
    def test_text_cells_preserved(self):
        data = _make_xlsx([["张三", "13800138000", "6222021000123456"]])
        self.assertEqual(xlsx_to_lines(data), ["张三\t13800138000\t6222021000123456"])

    def test_integer_no_sci(self):
        data = _make_xlsx([[1400, 20000, 150000]])
        self.assertEqual(xlsx_to_lines(data), ["1400\t20000\t150000"])

    def test_decimal(self):
        data = _make_xlsx([[100.5]])
        self.assertEqual(xlsx_to_lines(data), ["100.5"])

    def test_leading_zero_text(self):
        data = _make_xlsx([["00123"]])
        self.assertEqual(xlsx_to_lines(data), ["00123"])

    def test_empty_cells_kept_and_empty_rows_skipped(self):
        data = _make_xlsx([["a", None, "c"], [None, None, None], ["x"]])
        # 工作表有 3 列，最后一行也只有 3 列（后两列空）
        self.assertEqual(xlsx_to_lines(data), ["a\t\tc", "x\t\t"])

    def test_newline_in_cell_flattened(self):
        data = _make_xlsx([["a\nb", "c"]])
        self.assertEqual(xlsx_to_lines(data), ["a b\tc"])

    def test_bool(self):
        data = _make_xlsx([[True, False]])
        self.assertEqual(xlsx_to_lines(data), ["TRUE\tFALSE"])

    def test_date_only(self):
        data = _make_xlsx([[datetime.date(2024, 1, 5)]])
        self.assertEqual(xlsx_to_lines(data), ["2024-01-05"])

    def test_datetime(self):
        data = _make_xlsx([[datetime.datetime(2024, 1, 5, 13, 30, 45)]])
        self.assertEqual(xlsx_to_lines(data), ["2024-01-05 13:30:45"])

    def test_full_record_row(self):
        data = _make_xlsx(
            [["廖林培", "450421199010193538", 1400, "6217003400002221946", "18978486083"]]
        )
        self.assertEqual(
            xlsx_to_lines(data),
            ["廖林培\t450421199010193538\t1400\t6217003400002221946\t18978486083"],
        )

    def test_all_sheets_included(self):
        # 数据在第二个工作表，也应被读取到
        from openpyxl import Workbook

        wb = Workbook()
        ws1 = wb.active
        ws1.append(["这是第一页的表头"])
        ws2 = wb.create_sheet("第二页")
        ws2.append(["张三", "13800138000"])
        ws2.append(["李四", "13900139000"])
        buf = BytesIO()
        wb.save(buf)
        lines = xlsx_to_lines(buf.getvalue())
        self.assertIn("张三\t13800138000", lines)
        self.assertIn("李四\t13900139000", lines)
        self.assertIn("这是第一页的表头", lines)


if __name__ == "__main__":
    unittest.main(verbosity=2)
